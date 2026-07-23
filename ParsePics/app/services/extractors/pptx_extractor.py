"""Extract images from .pptx and replace them with textboxes."""

from __future__ import annotations

import logging
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.services.extractors.base import ExtractResult
from app.services.image_utils import bytes_to_png, image_label

logger = logging.getLogger(__name__)


def _is_picture(shape) -> bool:
    try:
        return shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    except NotImplementedError:
        # Some group/placeholder shapes raise NotImplementedError
        return False


def _collect_pictures(shapes) -> list:
    """Collect picture shapes, including those inside groups."""
    found = []
    for shape in list(shapes):
        if _is_picture(shape):
            found.append(shape)
            continue
        # Recurse into groups
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                found.extend(_collect_pictures(shape.shapes))
        except (NotImplementedError, AttributeError):
            continue
    return found


def process_pptx(source: Path, images_dir: Path, output_pptx: Path) -> ExtractResult:
    """
    Extract pictures, save PNGs, delete originals, add textboxes at same position.

    Returns ExtractResult with (label, png_path) entries and skip reasons.
    """
    prs = Presentation(str(source))
    result = ExtractResult()
    counter = 0

    for slide in prs.slides:
        pictures = _collect_pictures(slide.shapes)
        for shape in pictures:
            try:
                blob = shape.image.blob
            except Exception as exc:
                reason = f"pptx picture blob unavailable ({type(exc).__name__}: {exc})"
                result.skip_reasons.append(reason)
                logger.info("Skip image: %s", reason)
                continue

            left = shape.left
            top = shape.top
            width = shape.width
            height = shape.height

            counter += 1
            label = image_label(counter)
            png_path = images_dir / f"{label}.png"
            try:
                bytes_to_png(blob, png_path)
            except Exception as exc:
                counter -= 1
                reason = f"image convert failed ({type(exc).__name__}: {exc})"
                result.skip_reasons.append(reason)
                logger.warning("Skip image: %s", reason)
                continue

            # Remove original picture shape
            try:
                sp = shape._element
                sp.getparent().remove(sp)
            except Exception as exc:
                # If removal fails, still keep extracted image but skip textbox
                reason = f"picture removed incompletely ({type(exc).__name__}: {exc})"
                result.skip_reasons.append(reason)
                logger.warning("%s; keeping PNG %s", reason, label)
                result.entries.append((label, png_path))
                continue

            # Add textbox at the same position
            try:
                txbox = slide.shapes.add_textbox(left, top, width, height)
                tf = txbox.text_frame
                tf.word_wrap = True
                tf.text = label
            except Exception as exc:
                reason = f"textbox insert failed for {label} ({type(exc).__name__}: {exc})"
                result.skip_reasons.append(reason)
                logger.warning("%s; keeping PNG", reason)

            result.entries.append((label, png_path))

    output_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_pptx))
    return result
