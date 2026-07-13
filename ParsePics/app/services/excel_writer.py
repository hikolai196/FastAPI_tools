"""Build Excel mapping sheets with embedded PNGs."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter


def write_photo_excel(
    excel_path: Path,
    image_entries: list[tuple[str, Path]],
) -> Path:
    """
    Write {name, pic} Excel.

    image_entries: list of (label, png_path), e.g. ("img_001", Path(...))
    Even with zero images, writes a header-only workbook.
    """
    excel_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "photos"
    ws["A1"] = "name"
    ws["B1"] = "pic"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 28

    for row_idx, (name, png_path) in enumerate(image_entries, start=2):
        ws.cell(row=row_idx, column=1, value=name)
        if not png_path.exists():
            continue
        try:
            xl_img = XLImage(str(png_path))
            # Keep preview reasonably small
            max_w, max_h = 120, 120
            if xl_img.width and xl_img.height:
                scale = min(max_w / xl_img.width, max_h / xl_img.height, 1.0)
                xl_img.width = int(xl_img.width * scale)
                xl_img.height = int(xl_img.height * scale)
            ws.add_image(xl_img, f"{get_column_letter(2)}{row_idx}")
            ws.row_dimensions[row_idx].height = max(
                90, int((xl_img.height or 90) * 0.75)
            )
        except Exception:
            # Skip bad image; keep the name row
            ws.cell(row=row_idx, column=2, value="(image insert failed)")

    wb.save(excel_path)
    return excel_path
