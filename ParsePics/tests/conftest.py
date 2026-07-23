"""Build sample Office files with images for tests."""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pptx.util import Inches as PptxInches


def make_png_bytes(color: tuple[int, int, int] = (255, 0, 0), size: tuple[int, int] = (40, 40)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    return buf.getvalue()


def write_temp_png(path: Path, color: tuple[int, int, int] = (0, 128, 255)) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (48, 48), color).save(path, format="PNG")
    return path


def create_docx_with_images(path: Path, colors: list[tuple[int, int, int]]) -> Path:
    doc = Document()
    doc.add_paragraph("Before images")
    for i, color in enumerate(colors):
        png = path.parent / f"_tmp_img_{i}.png"
        write_temp_png(png, color)
        doc.add_picture(str(png), width=Inches(1.0))
        doc.add_paragraph(f"After image {i + 1}")
        png.unlink(missing_ok=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def create_docx_no_images(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("No images here")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def create_pptx_with_images(path: Path, colors: list[tuple[int, int, int]]) -> Path:
    prs = Presentation()
    blank = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(blank)
    for i, color in enumerate(colors):
        png = path.parent / f"_tmp_ppt_{i}.png"
        write_temp_png(png, color)
        slide.shapes.add_picture(
            str(png),
            PptxInches(0.5 + i * 1.5),
            PptxInches(1.0),
            width=PptxInches(1.2),
        )
        png.unlink(missing_ok=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def create_pptx_no_images(path: Path) -> Path:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_textbox(PptxInches(1), PptxInches(1), PptxInches(4), PptxInches(1)).text_frame.text = "hi"
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def excel_has_images(excel_path: Path) -> int:
    wb = load_workbook(excel_path)
    ws = wb.active
    return len(ws._images)
