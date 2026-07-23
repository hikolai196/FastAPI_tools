"""Unit / integration tests for ParsePics core + API."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.services.file_manager import ensure_dirs, prepare_output_dir, resolve_result_dir
from app.services.image_utils import bytes_to_png, image_label
from app.services.processor import process_file, process_files
from app.services.zip_writer import create_result_zip
from tests.conftest import (
    create_docx_no_images,
    create_docx_with_images,
    create_pptx_no_images,
    create_pptx_with_images,
    excel_has_images,
    make_png_bytes,
)


@pytest.fixture(autouse=True)
def _clean_result(tmp_path, monkeypatch):
    """Redirect result/temp under tmp to keep the repo clean."""
    import app.services.file_manager as fm

    result = tmp_path / "result"
    temp = tmp_path / "temp"
    result.mkdir()
    temp.mkdir()
    monkeypatch.setattr(fm, "RESULT_DIR", result)
    monkeypatch.setattr(fm, "TEMP_DIR", temp)
    monkeypatch.setattr(fm, "PROJECT_ROOT", tmp_path)
    ensure_dirs()
    yield result


def _result_dir_for(result_root: Path, base_name: str) -> Path:
    matches = sorted(result_root.glob(f"{base_name}_*"))
    dirs = [p for p in matches if p.is_dir()]
    assert dirs, f"No output dir for {base_name} under {result_root}"
    return dirs[-1]


def test_image_label():
    assert image_label(1) == "img_001"
    assert image_label(12) == "img_012"


def test_bytes_to_png(tmp_path: Path):
    dest = tmp_path / "out.png"
    bytes_to_png(make_png_bytes((10, 20, 30)), dest)
    assert dest.exists()
    assert dest.stat().st_size > 0


def test_unique_output_dirs_preserve_prior_runs(_clean_result, tmp_path: Path):
    src1 = create_docx_with_images(tmp_path / "report.docx", [(255, 0, 0)])
    r1 = process_file(src1)
    assert r1.status == "success"
    out1 = Path(_clean_result / Path(r1.output_dir).name)
    marker = out1 / "extracted_images" / "img_001.png"
    assert marker.exists()

    src2 = create_docx_with_images(tmp_path / "report.docx", [(0, 255, 0), (0, 0, 255)])
    r2 = process_file(src2)
    assert r2.status == "success"
    assert r1.output_dir != r2.output_dir
    assert marker.exists(), "first run must not be wiped"
    out2 = Path(_clean_result / Path(r2.output_dir).name)
    assert (out2 / "extracted_images" / "img_002.png").exists()
    assert r1.run_id and r2.run_id and r1.run_id != r2.run_id


def test_process_docx_with_images(_clean_result, tmp_path: Path):
    src = create_docx_with_images(
        tmp_path / "report.docx",
        [(255, 0, 0), (0, 255, 0), (0, 0, 255)],
    )
    result = process_file(src)
    assert result.status == "success"
    assert result.image_count == 3
    assert result.download_url and result.download_url.startswith("/download/")
    assert result.skipped_image_count == 0

    out = _result_dir_for(_clean_result, "report")
    assert (out / "extracted_images" / "img_001.png").exists()
    assert (out / "extracted_images" / "img_002.png").exists()
    assert (out / "extracted_images" / "img_003.png").exists()
    assert (out / "report_extracted.docx").exists()
    excel = out / "report_photoattacks.xlsx"
    assert excel.exists()
    assert excel_has_images(excel) == 3
    assert out.with_suffix(".zip").exists()

    wb = load_workbook(excel)
    ws = wb.active
    assert ws["A1"].value == "name"
    assert ws["B1"].value == "pic"
    assert ws["A2"].value == "img_001"
    assert ws["A4"].value == "img_003"


def test_process_docx_no_images(_clean_result, tmp_path: Path):
    src = create_docx_no_images(tmp_path / "empty.docx")
    result = process_file(src)
    assert result.status == "success"
    assert result.image_count == 0
    out = _result_dir_for(_clean_result, "empty")
    assert (out / "empty_extracted.docx").exists()
    excel = out / "empty_photoattacks.xlsx"
    assert excel.exists()
    assert excel_has_images(excel) == 0
    wb = load_workbook(excel)
    assert wb.active["A1"].value == "name"


def test_process_pptx_with_images(_clean_result, tmp_path: Path):
    src = create_pptx_with_images(
        tmp_path / "deck.pptx",
        [(255, 0, 0), (0, 255, 0)],
    )
    result = process_file(src)
    assert result.status == "success"
    assert result.image_count == 2

    out = _result_dir_for(_clean_result, "deck")
    assert (out / "extracted_images" / "img_001.png").exists()
    assert (out / "extracted_images" / "img_002.png").exists()
    assert (out / "deck_extracted.pptx").exists()
    excel = out / "deck_photoattacks.xlsx"
    assert excel.exists()
    assert excel_has_images(excel) == 2

    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(out / "deck_extracted.pptx"))
    texts = []
    pics = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            try:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    pics += 1
            except NotImplementedError:
                pass
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    assert pics == 0
    assert "img_001" in texts
    assert "img_002" in texts


def test_process_pptx_no_images(_clean_result, tmp_path: Path):
    src = create_pptx_no_images(tmp_path / "blank.pptx")
    result = process_file(src)
    assert result.status == "success"
    assert result.image_count == 0
    out = _result_dir_for(_clean_result, "blank")
    assert (out / "blank_extracted.pptx").exists()
    assert (out / "blank_photoattacks.xlsx").exists()


def test_process_files_isolates_failures(_clean_result, tmp_path: Path):
    good = create_docx_with_images(tmp_path / "good.docx", [(1, 2, 3)])
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a real docx")
    results = process_files([good, bad])
    assert len(results) == 2
    assert results[0].status == "success"
    assert results[1].status == "error"


def test_docx_skips_corrupt_image_blob(_clean_result, tmp_path: Path):
    """A drawing with unreadable blob is skipped; other images continue."""
    from docx import Document
    from docx.oxml.ns import qn

    from app.services.extractors.docx_extractor import process_docx

    src = create_docx_with_images(tmp_path / "mixed.docx", [(10, 20, 30), (40, 50, 60)])
    doc = Document(str(src))
    blip = next(doc.element.body.iter(qn("a:blip")))
    rid = blip.get(qn("r:embed"))
    part = doc.part.related_parts[rid]
    part._blob = b"not-an-image"

    corrupted = tmp_path / "corrupted.docx"
    doc.save(str(corrupted))

    out_dir = prepare_output_dir("mixed", run_id="testskip")
    extract = process_docx(
        corrupted,
        out_dir / "extracted_images",
        out_dir / "mixed_extracted.docx",
    )
    assert len(extract.entries) == 1
    assert extract.entries[0][0] == "img_001"
    assert len(extract.skip_reasons) >= 1
    assert any("convert failed" in r for r in extract.skip_reasons)


def test_create_result_zip(_clean_result, tmp_path: Path):
    out = prepare_output_dir("zipme", run_id="abc12345")
    (out / "extracted_images" / "img_001.png").write_bytes(make_png_bytes())
    (out / "note.txt").write_text("hi", encoding="utf-8")
    zpath = create_result_zip(out)
    assert zpath.exists()
    with zipfile.ZipFile(zpath) as zf:
        names = zf.namelist()
    assert any(n.endswith("img_001.png") for n in names)
    assert any(n.endswith("note.txt") for n in names)


def test_resolve_result_dir_rejects_traversal(_clean_result):
    assert resolve_result_dir("../etc") is None
    assert resolve_result_dir("missing_folder") is None
    out = prepare_output_dir("safe", run_id="deadbeef")
    assert resolve_result_dir(out.name) == out.resolve()


def test_health_and_process_api(_clean_result, tmp_path: Path):
    from app.api.main import app

    client = TestClient(app)
    assert client.get("/health").json() == {"status": "ok"}

    home = client.get("/")
    assert home.status_code == 200
    assert "ParsePics" in home.text
    assert client.get("/static/styles.css").status_code == 200
    assert client.get("/static/app.js").status_code == 200

    docx = create_docx_with_images(tmp_path / "api_doc.docx", [(9, 9, 9)])
    pptx = create_pptx_with_images(tmp_path / "api_ppt.pptx", [(8, 8, 8)])

    with docx.open("rb") as f1, pptx.open("rb") as f2:
        resp = client.post(
            "/process",
            files=[
                (
                    "files",
                    (
                        "api_doc.docx",
                        f1,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ),
                ),
                (
                    "files",
                    (
                        "api_ppt.pptx",
                        f2,
                        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    ),
                ),
            ],
        )

    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "success"
    assert len(body["results"]) == 2
    assert all(r["status"] == "success" for r in body["results"])
    assert body["results"][0]["image_count"] == 1
    assert body["results"][1]["image_count"] == 1
    assert body["results"][0]["download_url"].startswith("/download/")

    folder = Path(body["results"][0]["output_dir"]).name
    dl = client.get(f"/download/{folder}")
    assert dl.status_code == 200
    assert dl.headers["content-type"].startswith("application/zip")
    assert zipfile.is_zipfile(io.BytesIO(dl.content))


def test_api_download_not_found(_clean_result):
    from app.api.main import app

    client = TestClient(app)
    resp = client.get("/download/no_such_run")
    assert resp.status_code == 404


def test_api_rejects_unsupported(_clean_result, tmp_path: Path):
    from app.api.main import app

    client = TestClient(app)
    junk = tmp_path / "note.txt"
    junk.write_text("hello")
    with junk.open("rb") as f:
        resp = client.post(
            "/process",
            files=[("files", ("note.txt", f, "text/plain"))],
        )
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "error"
    assert body["results"][0]["status"] == "error"


def test_api_rejects_oversized_file(_clean_result, tmp_path: Path, monkeypatch):
    from app.api import main as api_main

    monkeypatch.setattr(api_main, "MAX_FILE_BYTES", 100)
    monkeypatch.setattr(api_main, "MAX_TOTAL_BYTES", 1000)

    client = TestClient(api_main.app)
    docx = create_docx_with_images(tmp_path / "big.docx", [(1, 1, 1)])
    assert docx.stat().st_size > 100
    with docx.open("rb") as f:
        resp = client.post(
            "/process",
            files=[
                (
                    "files",
                    (
                        "big.docx",
                        f,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ),
                )
            ],
        )
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "error"
    assert "MB limit" in body["results"][0]["error"]


def test_api_rejects_too_many_files(_clean_result, tmp_path: Path, monkeypatch):
    from app.api import main as api_main

    monkeypatch.setattr(api_main, "MAX_FILES_PER_REQUEST", 1)
    client = TestClient(api_main.app)

    a = create_docx_no_images(tmp_path / "a.docx")
    b = create_docx_no_images(tmp_path / "b.docx")
    with a.open("rb") as f1, b.open("rb") as f2:
        resp = client.post(
            "/process",
            files=[
                (
                    "files",
                    (
                        "a.docx",
                        f1,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ),
                ),
                (
                    "files",
                    (
                        "b.docx",
                        f2,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ),
                ),
            ],
        )
    assert resp.status_code == 400
    assert "Too many files" in resp.json()["detail"]
